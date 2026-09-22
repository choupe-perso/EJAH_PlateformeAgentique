# Créé par Cédric HOUPE.
# Usage ou reproduction à l'identique interdit sans autorisation.
"""PDF handler.

PDF has no reliable open-source primitive for *targeted* redaction: pypdf
(the pure-Python, BSD-licensed library used here) can only strip *all*
text from a page or every occurrence of a whole font, never "just this
word at this position" - and PyMuPDF, which does have that (via redaction
annotations), is AGPL-3.0, ruled out for this ADBI-proprietary tool.

Rather than fake a redaction - drawing a box with a token label over the
sensitive text while the original text stays technically present and
extractable underneath, which is not real anonymization - a PDF is
converted once into a plain new .docx (page text as paragraphs, embedded
images carried over) and then anonymized through the already
fully-reversible docx_handler pipeline. The output is a .docx, not a
.pdf: the original wording is genuinely gone from the file (not just
hidden), reversible the same way any other DOCX is - at the cost of the
PDF's original visual layout, fonts and exact formatting.

No deanonymize here: once converted, the output is a normal .docx and is
restored via docx_handler like any other one - there is no "anonymized
PDF" to restore back into a PDF.

Scanned pages (no extractable text layer) are out of scope: a warning is
printed and the page's images still go through the normal OCR/placeholder
pipeline (which usually catches a full-page scan, since that's typically
just one large embedded image), but no text-PII guarantee is made for
such a page.
"""
from __future__ import annotations

import io
import sys
import tempfile
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from app.tools.formats import docx_handler
from app.domain.span import Span
from app.domain.vault import Vault


def _pdf_to_docx(path: Path) -> Document:
    reader = PdfReader(str(path))
    document = Document()
    last_index = len(reader.pages) - 1
    for page_index, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if not text.strip():
            print(
                f"Avertissement : {path} page {page_index + 1} semble scannée "
                "(pas de texte extractible) ; la détection de texte n'est pas "
                "garantie sur cette page.",
                file=sys.stderr,
            )
        for line in text.splitlines():
            if line.strip():
                document.add_paragraph(line)

        for image in page.images:
            try:
                document.add_picture(io.BytesIO(image.data))
            except Exception:
                # A malformed/unsupported embedded image shouldn't abort
                # the whole conversion - skip it, the rest of the page
                # still gets processed normally.
                continue

        if page_index != last_index:
            document.add_paragraph()  # blank separator between pages

    return document


def _save_temp_docx(document: Document) -> Path:
    tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    tmp.close()
    document.save(tmp.name)
    return Path(tmp.name)


def inspect(path: Path) -> list[Span]:
    tmp_path = _save_temp_docx(_pdf_to_docx(path))
    try:
        return docx_handler.inspect(tmp_path)
    finally:
        tmp_path.unlink(missing_ok=True)


def anonymize(path: Path, vault: Vault, out_path: Path | None = None) -> tuple[Path, list[Span]]:
    tmp_path = _save_temp_docx(_pdf_to_docx(path))
    try:
        final_out = out_path if out_path is not None else path.with_suffix(".anonymise.docx")
        return docx_handler.anonymize(tmp_path, vault, final_out)
    finally:
        tmp_path.unlink(missing_ok=True)
