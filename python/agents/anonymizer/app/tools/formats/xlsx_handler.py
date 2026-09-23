# Créé par Cédric HOUPE.
# Usage ou reproduction à l'identique interdit sans autorisation.
"""XLSX handler: an Excel cell is a single atomic string, so - unlike
DOCX/PPTX - there's no run to split. Each non-empty string cell is
wrapped as a single-run "paragraph" and handed to `_run_splice`, which
already does exactly what's needed here: join with "\\n\\n" (so a span
never spans two cells - the paragraph-break truncation in
engine._trim_whitespace guarantees that, same as DOCX/PPTX), detect once
per sheet (helps cross-cell name/value propagation), then splice tokens
back in. Formula cells (`cell.data_type == "f"`) are never touched - the
formula source is never fed to detection. Whole embedded images are
swapped for a static placeholder when OCR finds PII in them (see
_ooxml_media).

Both directions are reversible. Text restoration reuses `_run_splice`
the same way. Images are trickier: unlike DOCX/PPTX, a floating Excel
image has no reliable alt-text field exposed by openpyxl to carry a vault
token. Instead, anonymize() records (media zip path -> image token) pairs
in a dedicated hidden worksheet (`_TOKENS_SHEET`), which deanonymize()
reads to know which vault entry restores which image, then removes -
the restored workbook has no trace of that bookkeeping sheet.

Known limitation to document: openpyxl doesn't losslessly round-trip
every advanced workbook feature (VBA macros unless `keep_vba=True`,
some chart types, pivot tables) - a general xlsx-fidelity risk on
save/reload, independent of the PII handling here.
"""
from __future__ import annotations

import io
from pathlib import Path

import openpyxl

from app.domain.engine import detect_only
from app.domain.entities import EntityType
from app.tools.formats import _ocr
from app.tools.formats._assets import placeholder_image_bytes
from app.tools.formats._ooxml_media import replace_zip_members
from app.tools.formats._run_splice import (
    detect_in_paragraphs,
    replace_in_paragraphs,
    restore_in_paragraphs,
)
from app.domain.span import Span
from app.domain.vault import Vault

_TOKENS_SHEET = "__anonymizer_tokens__"


class _CellText:
    """Adapts an openpyxl Cell to the RunLike (.text get/set) interface so
    _run_splice's offset-map algorithm can be reused directly - each cell
    plays the role of a single-run paragraph."""

    __slots__ = ("cell",)

    def __init__(self, cell) -> None:
        self.cell = cell

    @property
    def text(self) -> str:
        return self.cell.value

    @text.setter
    def text(self, value: str) -> None:
        self.cell.value = value


def _content_sheets(wb):
    return [ws for ws in wb.worksheets if ws.title != _TOKENS_SHEET]


def _string_cells(ws) -> list:
    cells = []
    for row in ws.iter_rows():
        for cell in row:
            if cell.data_type == "f":
                continue
            if isinstance(cell.value, str) and cell.value:
                cells.append(cell)
    return cells


def _sheet_paragraphs(ws) -> list[list[_CellText]]:
    return [[_CellText(cell)] for cell in _string_cells(ws)]


def _iter_images(wb):
    for ws in _content_sheets(wb):
        for image in getattr(ws, "_images", []):
            yield image


def _image_pii_spans(image_bytes: bytes) -> list[Span]:
    ocr_text = _ocr.image_text(image_bytes)
    if not ocr_text.strip():
        return []
    return detect_only(ocr_text)


def inspect(path: Path) -> list[Span]:
    wb = openpyxl.load_workbook(str(path))
    spans: list[Span] = []
    for ws in _content_sheets(wb):
        spans.extend(detect_in_paragraphs(_sheet_paragraphs(ws)))
    for image in _iter_images(wb):
        spans.extend(_image_pii_spans(image._data()))
    return spans


def anonymize(path: Path, vault: Vault, out_path: Path | None = None) -> tuple[Path, list[Span]]:
    wb = openpyxl.load_workbook(str(path))

    spans: list[Span] = []
    for ws in _content_sheets(wb):
        spans.extend(replace_in_paragraphs(_sheet_paragraphs(ws), vault))

    image_replacements: dict[str, bytes] = {}
    image_tokens: list[tuple[str, str]] = []
    for image in _iter_images(wb):
        image_bytes = image._data()
        # Image._data() reads (and, for png/jpeg/gif, closes) `image.ref`
        # in place - give it a fresh stream over the same bytes so
        # openpyxl's own save-time image writer can still read it.
        image.ref = io.BytesIO(image_bytes)
        img_spans = _image_pii_spans(image_bytes)
        if not img_spans:
            continue
        for s in img_spans:
            vault.tokenize(s.entity_type, s.text)
        spans.extend(img_spans)

        image_token = vault.tokenize_bytes(EntityType.IMAGE, image_bytes)
        member = str(image.path).lstrip("/")
        image_replacements[member] = placeholder_image_bytes()
        image_tokens.append((member, image_token))

    if image_tokens:
        tokens_ws = wb.create_sheet(_TOKENS_SHEET)
        tokens_ws.sheet_state = "hidden"
        for member, token in image_tokens:
            tokens_ws.append([member, token])

    out_path = out_path if out_path is not None else path.with_suffix(".anonymise.xlsx")
    wb.save(str(out_path))
    replace_zip_members(out_path, image_replacements)
    return out_path, spans


def deanonymize(path: Path, vault: Vault, out_path: Path | None = None) -> Path:
    wb = openpyxl.load_workbook(str(path))

    for ws in _content_sheets(wb):
        restore_in_paragraphs(_sheet_paragraphs(ws), vault)

    image_restorations: dict[str, bytes] = {}
    if _TOKENS_SHEET in wb.sheetnames:
        tokens_ws = wb[_TOKENS_SHEET]
        for member, token in tokens_ws.iter_rows(values_only=True):
            if not member or not token:
                continue
            original_bytes = vault.resolve_bytes(token)
            if original_bytes is not None:
                image_restorations[member] = original_bytes
        del wb[_TOKENS_SHEET]

    out_path = out_path if out_path is not None else path.with_suffix(".restored.xlsx")
    wb.save(str(out_path))
    replace_zip_members(out_path, image_restorations)
    return out_path
