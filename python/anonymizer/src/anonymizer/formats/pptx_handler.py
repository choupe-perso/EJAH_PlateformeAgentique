"""PPTX handler: same run-splice / whole-image-swap approach as
docx_handler (see _run_splice and _ooxml_media), applied to python-pptx's
model (shapes with text frames, tables, speaker notes). Both directions
are reversible: `deanonymize` restores the original text and the
original image bytes exactly, via the same vault used to anonymize.

The image token is stored in the picture's own `<p:cNvPr descr="...">`
attribute (standard alt-text field) - invisible in normal use, read back
by `deanonymize` to know which vault entry to restore each image from.

Known coverage gap, documented rather than silently swallowed:
- SmartArt and embedded chart text isn't exposed as ordinary text-frame
  runs by python-pptx, out of scope V1.
"""
from __future__ import annotations

import re
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from anonymizer.engine import detect_only
from anonymizer.entities import EntityType
from anonymizer.formats import _ocr
from anonymizer.formats._assets import placeholder_image_bytes
from anonymizer.formats._ooxml_media import replace_zip_members
from anonymizer.formats._run_splice import (
    detect_in_paragraphs,
    replace_in_paragraphs,
    restore_in_paragraphs,
)
from anonymizer.span import Span
from anonymizer.vault import Vault

_IMAGE_TOKEN_RE = re.compile(r"\[IMAGE_\d+\]")


def _iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)


def _all_paragraphs(prs: Presentation) -> list:
    paragraphs = []
    for slide in prs.slides:
        for shape in _iter_shapes(slide.shapes):
            if shape.has_text_frame:
                paragraphs.extend(shape.text_frame.paragraphs)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        paragraphs.extend(cell.text_frame.paragraphs)
        if slide.has_notes_slide:
            paragraphs.extend(slide.notes_slide.notes_text_frame.paragraphs)
    return paragraphs


def _find_c_nv_pr(blip):
    """Walk up from a <a:blip> to its enclosing <p:pic> and return that
    picture's <p:cNvPr> element (the alt-text/description field `descr`
    is stored on) - or None if the structure is unexpected (defensive
    only, never seen in practice)."""
    elm = blip
    while elm is not None:
        if elm.tag == qn("p:pic"):
            return elm.find(qn("p:nvPicPr") + "/" + qn("p:cNvPr"))
        elm = elm.getparent()
    return None


def _iter_image_parts(prs: Presentation):
    """Yields (image_part, c_nv_pr) for every embedded image across all
    slides (including inside groups) and their notes pages. `c_nv_pr` is
    the picture's alt-text element (or None if it couldn't be located) -
    used to tag/read the image's vault token for reversible restoration."""
    parts = set()
    for slide in prs.slides:
        parts.add(slide.part)
        if slide.has_notes_slide:
            parts.add(slide.notes_slide.part)
    for part in parts:
        for blip in part._element.iter(qn("a:blip")):
            rid = blip.get(qn("r:embed"))
            if not rid:
                continue
            rel = part.rels.get(rid)
            if rel is None or rel.is_external:
                continue
            yield rel.target_part, _find_c_nv_pr(blip)


def _image_pii_spans(image_bytes: bytes) -> list[Span]:
    ocr_text = _ocr.image_text(image_bytes)
    if not ocr_text.strip():
        return []
    return detect_only(ocr_text)


def inspect(path: Path) -> list[Span]:
    prs = Presentation(str(path))
    spans = detect_in_paragraphs([p.runs for p in _all_paragraphs(prs)])
    for image_part, _c_nv_pr in _iter_image_parts(prs):
        spans.extend(_image_pii_spans(image_part.blob))
    return spans


def anonymize(path: Path, vault: Vault, out_path: Path | None = None) -> tuple[Path, list[Span]]:
    prs = Presentation(str(path))

    paragraphs = [p.runs for p in _all_paragraphs(prs)]
    spans = replace_in_paragraphs(paragraphs, vault)

    image_replacements: dict[str, bytes] = {}
    for image_part, c_nv_pr in _iter_image_parts(prs):
        img_spans = _image_pii_spans(image_part.blob)
        if not img_spans:
            continue
        for s in img_spans:
            vault.tokenize(s.entity_type, s.text)
        spans.extend(img_spans)

        image_token = vault.tokenize_bytes(EntityType.IMAGE, image_part.blob)
        if c_nv_pr is not None:
            c_nv_pr.set("descr", image_token)
        member = str(image_part.partname).lstrip("/")
        image_replacements[member] = placeholder_image_bytes()

    out_path = out_path if out_path is not None else path.with_suffix(".anon.pptx")
    prs.save(str(out_path))
    replace_zip_members(out_path, image_replacements)
    return out_path, spans


def deanonymize(path: Path, vault: Vault, out_path: Path | None = None) -> Path:
    prs = Presentation(str(path))

    paragraphs = [p.runs for p in _all_paragraphs(prs)]
    restore_in_paragraphs(paragraphs, vault)

    image_restorations: dict[str, bytes] = {}
    for image_part, c_nv_pr in _iter_image_parts(prs):
        if c_nv_pr is None:
            continue
        descr = c_nv_pr.get("descr") or ""
        match = _IMAGE_TOKEN_RE.search(descr)
        if not match:
            continue
        original_bytes = vault.resolve_bytes(match.group(0))
        if original_bytes is None:
            continue
        member = str(image_part.partname).lstrip("/")
        image_restorations[member] = original_bytes
        c_nv_pr.set("descr", descr.replace(match.group(0), ""))

    out_path = out_path if out_path is not None else path.with_suffix(".restored.pptx")
    prs.save(str(out_path))
    replace_zip_members(out_path, image_restorations)
    return out_path
