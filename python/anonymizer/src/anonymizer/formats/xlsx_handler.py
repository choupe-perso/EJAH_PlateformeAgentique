"""XLSX handler: an Excel cell is a single atomic string, so - unlike
DOCX/PPTX - there's no run to split. Each non-empty string cell is
wrapped as a single-run "paragraph" and handed to `_run_splice`, which
already does exactly what's needed here: join with "\\n\\n" (so a span
never spans two cells - the paragraph-break truncation in
engine._trim_whitespace guarantees that, same as DOCX/PPTX), detect once
per sheet (helps cross-cell name/value propagation), then splice tokens
back in. Formula cells (`cell.data_type == "f"`) are never touched - the
formula source is never fed to detection. Whole embedded images are
swapped for a static placeholder when OCR finds PII in them (see
_ooxml_media). No deanonymize - restoring the original workbook isn't
implemented (see the project plan); the vault's token -> value table is
the audit trail.

Known limitation to document: openpyxl doesn't losslessly round-trip
every advanced workbook feature (VBA macros unless `keep_vba=True`,
some chart types, pivot tables) - a general xlsx-fidelity risk on
save/reload, independent of the PII handling here.
"""
from __future__ import annotations

import io
from pathlib import Path

import openpyxl

from anonymizer.engine import detect_only
from anonymizer.formats import _ocr
from anonymizer.formats._assets import placeholder_image_bytes
from anonymizer.formats._ooxml_media import replace_zip_members
from anonymizer.formats._run_splice import detect_in_paragraphs, replace_in_paragraphs
from anonymizer.span import Span
from anonymizer.vault import Vault


class _CellText:
    """Adapts an openpyxl Cell to the RunLike (.text get/set) interface so
    _run_splice's offset-map algorithm can be reused directly - each cell
    plays the role of a single-run paragraph."""

    __slots__ = ("cell",)

    def __init__(self, cell) -> None:
        self.cell = cell

    @property
    def text(self) -> str:
        return self.cell.value

    @text.setter
    def text(self, value: str) -> None:
        self.cell.value = value


def _string_cells(ws) -> list:
    cells = []
    for row in ws.iter_rows():
        for cell in row:
            if cell.data_type == "f":
                continue
            if isinstance(cell.value, str) and cell.value:
                cells.append(cell)
    return cells


def _sheet_paragraphs(ws) -> list[list[_CellText]]:
    return [[_CellText(cell)] for cell in _string_cells(ws)]


def _iter_images(wb):
    for ws in wb.worksheets:
        for image in getattr(ws, "_images", []):
            yield image


def _image_pii_spans(image_bytes: bytes) -> list[Span]:
    ocr_text = _ocr.image_text(image_bytes)
    if not ocr_text.strip():
        return []
    return detect_only(ocr_text)


def inspect(path: Path, avec_images: bool = True) -> list[Span]:
    wb = openpyxl.load_workbook(str(path))
    spans: list[Span] = []
    for ws in wb.worksheets:
        spans.extend(detect_in_paragraphs(_sheet_paragraphs(ws)))
    if avec_images:
        for image in _iter_images(wb):
            spans.extend(_image_pii_spans(image._data()))
    return spans


def anonymize(
    path: Path, vault: Vault, out_path: Path | None = None, avec_images: bool = True
) -> tuple[Path, list[Span]]:
    wb = openpyxl.load_workbook(str(path))

    spans: list[Span] = []
    for ws in wb.worksheets:
        spans.extend(replace_in_paragraphs(_sheet_paragraphs(ws), vault))

    image_replacements: dict[str, bytes] = {}
    if avec_images:
        for image in _iter_images(wb):
            image_bytes = image._data()
            # Image._data() reads (and, for png/jpeg/gif, closes) `image.ref`
            # in place - give it a fresh stream over the same bytes so
            # openpyxl's own save-time image writer can still read it.
            image.ref = io.BytesIO(image_bytes)
            img_spans = _image_pii_spans(image_bytes)
            if not img_spans:
                continue
            for s in img_spans:
                vault.tokenize(s.entity_type, s.text)
            spans.extend(img_spans)
            member = str(image.path).lstrip("/")
            image_replacements[member] = placeholder_image_bytes()

    out_path = out_path if out_path is not None else path.with_suffix(".anon.xlsx")
    wb.save(str(out_path))
    replace_zip_members(out_path, image_replacements)
    return out_path, spans
